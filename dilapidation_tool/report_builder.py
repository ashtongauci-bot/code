# report_builder.py - Builds the Word document matching the IIE template format

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
from config import PROJECT
from generate_rating_images import ensure_rating_images

TNR = "Times New Roman"
APTOS = "Aptos"


def set_run(run, size=None, bold=None, italic=None, underline=None, font_name=TNR):
    if font_name:
        run.font.name = font_name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def _setup_document(doc):
    """A4 page size, spec margins, en-AU language — applied to fresh documents only."""
    for section in doc.sections:
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
        section.top_margin = Twips(1134)
        section.bottom_margin = Twips(1134)
        section.left_margin = Twips(851)
        section.right_margin = Twips(851)
        section.header_distance = Twips(567)
        section.footer_distance = Twips(284)
    try:
        rPr = doc.styles["Normal"].element.get_or_add_rPr()
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "en-AU")
        rPr.append(lang)
    except Exception:
        pass


def _add_page_number_footer(doc):
    """Add centred [ N ] page-number footer to every section."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(fp._p.findall(qn("w:r"))):
            fp._p.remove(r)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def _append_run(tag, text=None, fld_type=None):
            r = OxmlElement("w:r")
            if tag == "w:fldChar":
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), fld_type)
            elif tag == "w:instrText":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            else:
                el = OxmlElement(tag)
                if text:
                    el.set(qn("xml:space"), "preserve")
                    el.text = text
            r.append(el)
            fp._p.append(r)

        _append_run("w:t", text="[ ")
        _append_run("w:fldChar", fld_type="begin")
        _append_run("w:instrText", text=" PAGE ")
        _append_run("w:fldChar", fld_type="separate")
        _append_run("w:t", text="1")
        _append_run("w:fldChar", fld_type="end")
        _append_run("w:t", text=" ]")


def add_cover_page(doc, cover_photo: str = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PRE-CONSTRUCTION DILAPIDATION REPORT")
    set_run(run, size=18, bold=True, underline=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("At"), size=12, font_name=APTOS)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Surrounding Council Assets"), size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Due to development at:"), size=12, font_name=APTOS)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(PROJECT["address"]), size=18, bold=True)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(f"Prepared For: {PROJECT['client']}"), size=12, bold=True)

    doc.add_paragraph()

    if cover_photo and Path(cover_photo).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(cover_photo), width=Inches(5.0))
    else:
        doc.add_paragraph()
        doc.add_paragraph()

    doc.add_paragraph()

    for label, value in [("Date:", PROJECT["report_date"]), ("Ref:", PROJECT["ref"])]:
        p = doc.add_paragraph()
        set_run(p.add_run(f"{label}\t{value}"), size=12, font_name=APTOS)

    doc.add_page_break()


def _short_address(full: str) -> str:
    parts = [p.strip() for p in full.split(",")]
    return ", ".join(parts[:2])


def _make_borderless_table(doc, num_rows, num_cols, col_widths_dxa):
    """Create a borderless table with explicit column and table widths."""
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = "Normal Table"
    total_width = sum(col_widths_dxa)

    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total_width))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    tblBorders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        tblBorders.append(el)
    tblPr.append(tblBorders)

    return table


def _set_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        tcBorders.append(el)
    tcPr.append(tcBorders)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def _caption_para(cell, photo_number: int, description: str):
    """Caption paragraph: TNR 9pt bold label + plain description, #0E2841, centred."""
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = cap._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "200")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)

    colour = RGBColor(0x0E, 0x28, 0x41)
    r1 = cap.add_run(f"Photograph {photo_number}:")
    set_run(r1, size=9, bold=True)
    r1.font.color.rgb = colour

    r2 = cap.add_run(f" {description}")
    set_run(r2, size=9, bold=False)
    r2.font.color.rgb = colour


def add_metadata_block(doc, rows: list[tuple[str, str]]):
    """Borderless 2-column table for report metadata."""
    label_w, value_w = 2304, 7624
    table = _make_borderless_table(doc, len(rows), 2, [label_w, value_w])

    for i, (label, value) in enumerate(rows):
        row = table.rows[i]
        lc, vc = row.cells[0], row.cells[1]
        _set_cell_width(lc, label_w)
        _set_cell_width(vc, value_w)
        set_run(lc.paragraphs[0].add_run(label), size=12, bold=True)
        set_run(vc.paragraphs[0].add_run(value), size=12, font_name=APTOS)

    doc.add_paragraph()


def add_report_metadata(doc):
    short_addr = _short_address(PROJECT["address"])
    add_metadata_block(doc, [
        ("Name:",               f"Pre-Construction Dilapidation Report – {short_addr}"),
        ("Date of Inspection:", PROJECT["inspection_date"]),
        ("To:",                 PROJECT["client"]),
    ])


def _toc_entry_para(doc, text, page_str):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9000")
    tabs.append(tab)
    pPr.append(tabs)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "480")
    spacing.set(qn("w:after"), "0")
    pPr.append(spacing)
    set_run(p.add_run(f"{text}\t{page_str}"), size=12, italic=True)
    return p


def add_contents(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("TABLE OF CONTENTS"), size=12, bold=True, underline=True)
    doc.add_paragraph()

    p_toc = doc.add_paragraph()
    r1 = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    fc.set(qn("w:dirty"), "true")
    r1.append(fc)
    p_toc._p.append(r1)

    r2 = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-1" \\h \\z \\u '
    r2.append(instr)
    p_toc._p.append(r2)

    r3 = OxmlElement("w:r")
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "separate")
    r3.append(fc3)
    p_toc._p.append(r3)

    for text, page in [
        ("1.0 PREAMBLE", "3"), ("2.0 INTRODUCTION", "3"),
        ("3.0 EXISTING CONDITIONS", "6"), ("4.0 CONCLUSION", "66"),
    ]:
        _toc_entry_para(doc, text, page)

    r4 = OxmlElement("w:r")
    fc4 = OxmlElement("w:fldChar")
    fc4.set(qn("w:fldCharType"), "end")
    r4.append(fc4)
    doc.paragraphs[-1]._p.append(r4)

    doc.add_page_break()


def add_section_heading(doc, text):
    """Heading 1 — TNR bold black, spacing before 360 after 80."""
    p = doc.add_paragraph(style="Heading 1")
    pPr = p._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "360")
    spacing.set(qn("w:after"), "80")
    pPr.append(spacing)
    run = p.add_run(text)
    set_run(run, bold=True)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_sub_label(doc, text):
    p = doc.add_paragraph()
    set_run(p.add_run(text), size=9, bold=True)
    return p


def add_body(doc, text):
    """Body paragraph — Aptos 12pt justified."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p.add_run(text), size=12, font_name=APTOS)
    return p


def add_mixed_para(doc, parts):
    """Mixed bold/plain paragraph — Aptos 12pt justified."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, bold in parts:
        set_run(p.add_run(text), size=12, bold=(True if bold else None), font_name=APTOS)
    return p


def add_bullet(doc, text):
    try:
        p = doc.add_paragraph(style="List Bullet")
    except Exception:
        p = doc.add_paragraph()
    set_run(p.add_run(text), size=12, font_name=APTOS)
    return p


def add_preamble(doc):
    add_section_heading(doc, "1.0 PREAMBLE")
    add_mixed_para(doc, [
        ("This pre-construction dilapidation report is based on visual inspection only. "
         "The purpose of this report is to provide a photographic record of the ", False),
        (f"Council Assets along {PROJECT['streets_inspected']}", True),
        (". The council assets include roads and footpaths within the zone of "
         "influence of the proposed construction site at ", False),
        (PROJECT["address"], True),
        (".", False),
    ])
    add_body(doc, (
        "This report also gives a brief descriptive record of any defects noted on the date of our inspection. "
        "The inspection included all site features and accessible areas of the council assets as photographed "
        "and identified within this report. Photos show items of note, such as cracks, as well as some overviews."
    ))
    add_body(doc, (
        f"A total of {PROJECT['total_photos']} photos was taken during our inspection ({PROJECT['inspection_date']}). "
        f"A full set of photos can be downloaded via the following link: {PROJECT['photo_link']}"
    ))
    add_body(doc, (
        "This report is not a structural or civil engineering report. It is the property owner's responsibility "
        "to seek further structural engineering advice on any defective elements which have been noted in this report."
    ))
    doc.add_paragraph()


def add_map_figure(doc, image_path: str, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    try:
        run.add_picture(image_path, width=Inches(6.0))
    except Exception:
        p.add_run(f"[Map image not found: {image_path}]")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(cap.add_run(caption), size=12, bold=True, italic=True, font_name=APTOS)


def add_introduction(doc, map_paths: dict = None):
    map_paths = map_paths or {}
    ensure_rating_images(Path(__file__).parent)
    add_section_heading(doc, "2.0 INTRODUCTION")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p.add_run("The inspection focused conditions to the council assets that surround "), size=12, font_name=APTOS)
    set_run(p.add_run(PROJECT["address"]), size=12, bold=True, font_name=APTOS)
    set_run(p.add_run(
        ". The 50m inspection corridor extending either side of the subject address "
        "is illustrated in Figure 3."
    ), size=12, font_name=APTOS)

    doc.add_paragraph()
    if map_paths.get("figure3"):
        add_map_figure(doc, map_paths["figure3"], "Figure 3 – 50m Inspection Corridor (Not to Scale)")
    else:
        p = doc.add_paragraph()
        set_run(p.add_run("Figure 3 – 50m Inspection Corridor (Not to Scale)"), size=12, bold=True, italic=True, font_name=APTOS)
    doc.add_paragraph()

    add_body(doc, (
        "The areas inspected include all road pavement surfaces, pathways, stairs, concrete footpaths, "
        "grass, kerb and gutters, vehicular crossings, in-ground service pits, street trees and signs within "
        "the vicinity of the subject development. Specifically, when discussing the council assets herein, "
        "we discuss them in the following sub-categories:"
    ))

    add_bullet(doc, "Surrounding Roads and Pathways (Building Side)")
    add_bullet(doc, "Surrounding Roads and Pathways (Opposite Side)")

    doc.add_paragraph()
    add_body(doc, "Description of terms in the report (based on visual observations) are:")

    for term, definition in [
        ("Good –", "Items do not appear to have any defects and are in good condition."),
        ("Fair –", "Item is in reasonable condition for its age and may have some minor defect(s)."),
        ("Poor –", "Indicates generally that a defect is beyond minor and further structural advice may be required."),
    ]:
        p = doc.add_paragraph()
        set_run(p.add_run(f"{term}    "), size=12, bold=True, font_name=APTOS)
        set_run(p.add_run(definition), size=12, font_name=APTOS)

    doc.add_paragraph()

    rating_graph = Path(__file__).parent / "pavement_rating_graph.png"
    if rating_graph.exists():
        add_map_figure(doc, str(rating_graph), "Figure 4 – Graph Depicting Pavement Rating System")
    else:
        p = doc.add_paragraph()
        set_run(p.add_run("Figure 4 – Graph Depicting Pavement Rating System"), size=12, bold=True, italic=True, font_name=APTOS)

    doc.add_paragraph()

    rating_table = Path(__file__).parent / "pavement_rating_table.png"
    if rating_table.exists():
        add_map_figure(doc, str(rating_table), "Figure 5 – Table Describing Pavement Rating System in More Detail")
    else:
        p = doc.add_paragraph()
        set_run(p.add_run("Figure 5 – Table Describing Pavement Rating System in More Detail"), size=12, bold=True, italic=True, font_name=APTOS)

    doc.add_page_break()


def add_existing_conditions_intro(doc):
    add_section_heading(doc, "3.0 EXISTING CONDITIONS")
    add_mixed_para(doc, [
        ("The inspection covered all council-managed roads and footpaths along ", False),
        (PROJECT["streets_inspected"], True),
        (" within the zone of influence of the proposed development at ", False),
        (PROJECT["address"], True),
        (". The roads and pathways were found to be in fair to poor condition with some cracks present "
         "typical for the age of the infrastructure. Photos and description of condition can be found overleaf.", False),
    ])
    doc.add_paragraph()
    add_sub_label(doc, "SURROUNDING ROAD AND PATHWAYS")
    doc.add_paragraph()


def add_photo_table_entry(doc, photo: dict):
    """Each photo in a 2-row single-column table: image row + caption row."""
    photo_path = Path(photo["path"])
    content_width = 10204  # 11906 - 2×851 DXA

    table = _make_borderless_table(doc, 2, 1, [content_width])

    # Centre table on page
    tblPr = table._tbl.find(qn("w:tblPr"))
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblPr.append(jc)

    img_cell = table.cell(0, 0)
    _set_cell_width(img_cell, content_width)
    img_para = img_cell.paragraphs[0]
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_para.add_run()
    if photo_path.exists():
        try:
            run.add_picture(str(photo_path), width=Inches(4.73))
        except Exception:
            img_para.add_run(f"[Photo: {photo_path.name}]")
    else:
        img_para.add_run(f"[Photo not found: {photo_path.name}]")

    cap_cell = table.cell(1, 0)
    _set_cell_width(cap_cell, content_width)
    _caption_para(cap_cell, photo["number"], photo["description"])

    doc.add_paragraph()


def add_direction_heading(doc, section_name: str):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(section_name), size=9, bold=True)
    doc.add_paragraph()


def _add_signature_table(doc):
    """Sign-off block using tab stops (not a table)."""
    base = Path(__file__).parent
    inspector_sig = base / "inspector_signature.png"
    reviewer_sig = base / "reviewer_signature.png"
    inspector_full = f"{PROJECT.get('inspector_title', '')} {PROJECT['inspector_name']}".strip()
    reviewer_full = PROJECT.get("reviewer_name", "")
    inspector_quals = PROJECT.get("inspector_quals", "").split("\n")
    reviewer_quals = PROJECT.get("reviewer_quals", "").split("\n")

    def _tabbed_para():
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), "4800")
        tabs.append(tab)
        pPr.append(tabs)
        return p

    p = _tabbed_para()
    set_run(p.add_run("Yours faithfully,"), size=12, font_name=APTOS)
    p.add_run("\t")
    set_run(p.add_run("Reviewed by,"), size=12, font_name=APTOS)

    p = _tabbed_para()
    r_left = p.add_run()
    if inspector_sig.exists():
        r_left.add_picture(str(inspector_sig), height=Inches(0.6))
    p.add_run("\t")
    r_right = p.add_run()
    if reviewer_sig.exists():
        r_right.add_picture(str(reviewer_sig), height=Inches(0.6))

    if not inspector_sig.exists() and not reviewer_sig.exists():
        for _ in range(3):
            doc.add_paragraph()

    p = _tabbed_para()
    set_run(p.add_run(inspector_full), size=12, bold=True, font_name=APTOS)
    p.add_run("\t")
    set_run(p.add_run(reviewer_full), size=12, bold=True, font_name=APTOS)

    for i in range(max(len(inspector_quals), len(reviewer_quals))):
        p = _tabbed_para()
        left = inspector_quals[i] if i < len(inspector_quals) else ""
        right = reviewer_quals[i] if i < len(reviewer_quals) else ""
        set_run(p.add_run(left), size=12, italic=True, font_name=APTOS)
        p.add_run("\t")
        set_run(p.add_run(right), size=12, italic=True, font_name=APTOS)

    doc.add_paragraph()
    p = doc.add_paragraph()
    set_run(p.add_run(f"For, and on behalf of, {PROJECT['company']}."), size=12, font_name=APTOS)


def add_conclusion(doc):
    doc.add_page_break()
    add_section_heading(doc, "4.0 CONCLUSION")
    add_mixed_para(doc, [
        ("The inspection covered all council-managed roads and footpaths along ", False),
        (PROJECT["streets_inspected"], True),
        (" within the zone of influence of the proposed development at ", False),
        (PROJECT["address"], True),
        (". The roads and pathways were found to be in reasonable to poor condition with some cracks present "
         "typical for the age of the infrastructure. No signs of significant structural distress were observed.", False),
    ])
    add_body(doc, (
        "The roads and footpaths were inspected on both sides of the roads. Across the road and footpaths there "
        "was cracking present which showed typical wear of council assets. These items have been documented in "
        "the photographic record above for future reference."
    ))
    add_body(doc, (
        f"A total of {PROJECT['total_photos']} photos was taken during our inspection ({PROJECT['inspection_date']}). "
        f"A full set of photos can be downloaded via the following link: {PROJECT['photo_link']}"
    ))
    add_body(doc, (
        "I am an appropriately qualified and person competent in this area. I possess indemnity insurance to "
        "the satisfaction of the client. We trust that this dilapidation report meets your requirements."
    ))
    doc.add_paragraph()
    doc.add_paragraph()
    _add_signature_table(doc)


def build_report(all_photos: list[dict], output_path: str, template_path: str = None, map_paths: dict = None):
    if template_path and Path(template_path).exists():
        doc = Document(template_path)
        body = doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        for child in list(body):
            body.remove(child)
        if sectPr is not None:
            body.append(sectPr)
        print(f"Using template: {Path(template_path).name}")
    else:
        doc = Document()
        _setup_document(doc)

    _add_page_number_footer(doc)

    add_cover_page(doc, cover_photo=map_paths.get("cover") if map_paths else None)
    add_contents(doc)
    add_report_metadata(doc)
    add_preamble(doc)
    add_introduction(doc, map_paths=map_paths)
    add_existing_conditions_intro(doc)

    sections_seen = []
    for photo in all_photos:
        section = photo["section"]
        if not sections_seen or sections_seen[-1] != section:
            sections_seen.append(section)
            add_direction_heading(doc, section)
        add_photo_table_entry(doc, photo)

    add_conclusion(doc)
    doc.save(output_path)
    print(f"\nReport saved to: {output_path}")
