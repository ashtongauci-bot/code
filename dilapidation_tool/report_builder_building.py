# report_builder_building.py - Builds building dilapidation reports

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
from config import PROJECT
from report_builder import (
    TNR, APTOS, set_run,
    _short_address, _setup_document, _add_page_number_footer,
    _make_borderless_table, _set_cell_width, _caption_para,
    _add_signature_table,
)


def add_body(doc, text, size=12, bold=None, italic=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold, italic=italic, font_name=APTOS)
    return p


def add_mixed_para(doc, parts):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, bold in parts:
        set_run(p.add_run(text), size=12, bold=(True if bold else None), font_name=APTOS)
    return p


def add_section_heading(doc, text):
    """Heading 1 — TNR bold black, spacing before 360 after 80."""
    p = doc.add_paragraph(style="Heading 1")
    pPr = p._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "360")
    spacing.set(qn("w:after"), "80")
    pPr.append(spacing)
    rPr = p.runs[0]._r.get_or_add_rPr() if p.runs else None
    run = p.add_run(text)
    for tag, attrs in [
        ("w:rFonts", {"w:ascii": TNR, "w:hAnsi": TNR, "w:cs": TNR}),
        ("w:b",      {}),
        ("w:color",  {"w:val": "000000"}),
        ("w:sz",     {"w:val": "24"}),
        ("w:u",      {"w:val": "single"}),
    ]:
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        run._r.get_or_add_rPr().append(el)
    return p


def add_heading2(doc, text):
    """Heading 2 — 16pt, colour #0F4761, spacing before 160 after 80."""
    p = doc.add_paragraph(style="Heading 2")
    pPr = p._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "160")
    spacing.set(qn("w:after"), "80")
    pPr.append(spacing)
    run = p.add_run(text)
    set_run(run, size=16)
    run.font.color.rgb = RGBColor(0x0F, 0x47, 0x61)
    return p


def add_appendix_label(doc, text):
    p = doc.add_paragraph()
    set_run(p.add_run(text), size=12, bold=True, font_name=APTOS)
    return p


def add_facade_label(doc, text):
    p = doc.add_paragraph()
    set_run(p.add_run(text), size=12, bold=True, font_name=APTOS)
    return p


def add_cover_page(doc, cover_photo: str = None):
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("PRE-CONSTRUCTION DILAPIDATION REPORT"), size=18, bold=True, underline=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("At"), size=12, font_name=APTOS)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(PROJECT["address"]), size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Due to development at:"), size=12, font_name=APTOS)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(PROJECT.get("development_address", PROJECT["address"])), size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(f"Prepared For: {PROJECT['client']}"), size=12, bold=True, font_name=APTOS)

    doc.add_paragraph()

    manual_cover = Path(__file__).parent / "photos" / "cover.jpg"
    photo_to_use = cover_photo or (str(manual_cover) if manual_cover.exists() else None)
    if photo_to_use and Path(photo_to_use).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(photo_to_use), width=Inches(5.5))
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("[Cover photo – place cover.jpg in photos folder]")
        set_run(run, size=11, italic=True, font_name=APTOS)

    doc.add_paragraph()

    for label, value in [("Date:", PROJECT["report_date"]), ("Ref:", PROJECT["ref"])]:
        p = doc.add_paragraph()
        set_run(p.add_run(f"{label}\t{value}"), size=12, font_name=APTOS)

    doc.add_page_break()


def _toc_entry_para(doc, text, page_str):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9000")
    tabs.append(tab)
    pPr.append(tabs)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "480")
    spacing.set(qn("w:after"), "0")
    pPr.append(spacing)
    set_run(p.add_run(f"{text}\t{page_str}"), size=12, italic=True)
    return p


def add_contents(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("TABLE OF CONTENTS"), size=12, bold=True, underline=True)
    doc.add_paragraph()

    p_toc = doc.add_paragraph()
    r1 = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    fc.set(qn("w:dirty"), "true")
    r1.append(fc)
    p_toc._p.append(r1)

    r2 = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-1" \\h \\z \\u '
    r2.append(instr)
    p_toc._p.append(r2)

    r3 = OxmlElement("w:r")
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "separate")
    r3.append(fc3)
    p_toc._p.append(r3)

    for text, page in [
        ("1.0 PREAMBLE", "3"), ("2.0 INTRODUCTION", "3"),
        ("3.0 EXISTING CONDITIONS", "6"), ("4.0 CONCLUSION", "66"),
        ("APPENDIX A – EXTERNAL FACADES", ""), ("APPENDIX B – INTERNAL AREAS", ""),
    ]:
        _toc_entry_para(doc, text, page)

    r4 = OxmlElement("w:r")
    fc4 = OxmlElement("w:fldChar")
    fc4.set(qn("w:fldCharType"), "end")
    r4.append(fc4)
    doc.paragraphs[-1]._p.append(r4)

    doc.add_page_break()


def add_report_metadata(doc):
    from report_builder import add_metadata_block
    short_addr = _short_address(PROJECT["address"])
    add_metadata_block(doc, [
        ("Name:",               f"Pre-Construction Dilapidation Report – {short_addr}"),
        ("Date of Inspection:", PROJECT["inspection_date"]),
        ("To:",                 PROJECT["client"]),
    ])


def add_preamble(doc):
    add_section_heading(doc, "1.0 PREAMBLE")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p.add_run(
        "This pre-construction dilapidation report is based on visual inspection only. "
        "The purpose of this report is to provide a photographic record of the "
        "existing internal and external conditions of "
    ), size=12, font_name=APTOS)
    set_run(p.add_run(f"{PROJECT['address']}."), size=12, bold=True, font_name=APTOS)

    prop_desc = PROJECT.get("property_description", "")
    if prop_desc:
        set_run(p.add_run(f" {prop_desc}"), size=12, font_name=APTOS)

    dev_addr = PROJECT.get("development_address", "")
    if dev_addr:
        set_run(p.add_run(" This property is located adjacent to the proposed development at "), size=12, font_name=APTOS)
        set_run(p.add_run(f"{dev_addr}."), size=12, bold=True, font_name=APTOS)

    set_run(p.add_run(
        " This report also gives a brief descriptive record of any defects noted on the "
        "date of our inspection. The inspection included all site features and accessible "
        "areas of the property as photographed and identified within this report. "
        "Photos show items of note, such as cracks, as well as some overviews."
    ), size=12, font_name=APTOS)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p2.add_run(
        f"A full set of photos taken during our inspection ({PROJECT['inspection_date']}) "
        "can be provided upon request."
    ), size=12, font_name=APTOS)

    add_body(doc, (
        "This report is not a structural or civil engineering report. It is the property owner's "
        "responsibility to seek further structural engineering advice on any defective elements "
        "which have been noted in this report."
    ))
    doc.add_paragraph()


def add_introduction(doc, map_paths: dict = None):
    map_paths = map_paths or {}
    add_section_heading(doc, "2.0 INTRODUCTION")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p.add_run("The inspection focused on the internal and external conditions to "), size=12, font_name=APTOS)
    set_run(p.add_run(f"{PROJECT['address']}."), size=12, bold=True, font_name=APTOS)
    doc.add_paragraph()

    add_body(doc, (
        "The areas inspected at this property include all external building facades and external "
        "site features. The internals for the property were also inspected as access was provided. "
        "Specifically, when discussing the pre-construction condition of the building in this report, "
        "we have structured this report in the following sub-categories:"
    ))

    for item in ["External Facades", "Internals"]:
        try:
            p = doc.add_paragraph(style="List Bullet")
        except Exception:
            p = doc.add_paragraph()
        set_run(p.add_run(item), size=12, font_name=APTOS)

    doc.add_paragraph()
    add_body(doc, (
        "When describing the condition of an element in this report (Good, Fair or Poor), these are "
        "visual observation-based opinions and are based on the following definitions:"
    ))

    for term, definition in [
        ("Good –", "Items do not appear to have any defects and are in good condition."),
        ("Fair –", "Item is in reasonable condition for its age and may have some minor defect(s)."),
        ("Poor –", "Indicates generally that a defect is beyond minor and further structural advice is required."),
    ]:
        p = doc.add_paragraph()
        set_run(p.add_run(f"{term}    "), size=12, bold=True, font_name=APTOS)
        set_run(p.add_run(definition), size=12, font_name=APTOS)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(p.add_run("Also, when categorising cracking in this report, we rely on the definitions defined in "), size=12, font_name=APTOS)
    set_run(p.add_run(
        "AS2870-2011, Appendix C (Page 72), Table C1 and the NSW Guide to Standards and Tolerances, 2017 (NSWGST)"
    ), size=12, italic=True, font_name=APTOS)
    set_run(p.add_run(
        " as per extracts in Figures 3 & 4 below. Although, this extract classifies damage due to "
        "foundation movements (particularly cracking with reference to walls) we still believe it to "
        "be useful in categorizing all cracking defects identified at our dilapidation survey inspection. "
        "Cracks in this report are therefore categorised as follows:"
    ), size=12, font_name=APTOS)

    doc.add_paragraph()

    for fig_file, caption in [
        ("rating_graph.png",   "Figure 3 – AS2870 Classification of Damage Due to Foundation Movements"),
        ("rating_table.png",   "Figure 4 – Table 3.02 Damage to Walls Caused by Movement of Slabs and Footings"),
    ]:
        fig_path = Path(__file__).parent / fig_file
        if fig_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(fig_path), width=Inches(5.49))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(cap.add_run(caption), size=12, bold=True, italic=True, font_name=APTOS)
            doc.add_paragraph()
        else:
            print(f"  Warning: {fig_file} not found – figure skipped")

    add_body(doc, (
        "As AS2870 and the NSWGST extracts suggest, Category 3 cracking is considered structural, "
        "and therefore further structural engineering advice on such defects should be sourced. "
        "Figure 5 is another extract from the NSWGST which further validates this opinion:"
    ))
    doc.add_paragraph()

    fig5 = Path(__file__).parent / "rating_extract.png"
    if fig5.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(fig5), width=Inches(5.49))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cap.add_run("Figure 5 – NSW Guide to Standards and Tolerances, 2017 (NSWGST) Extract"), size=12, bold=True, italic=True, font_name=APTOS)
    else:
        print("  Warning: rating_extract.png not found – Figure 5 skipped")

    doc.add_page_break()


def add_existing_conditions_intro(doc):
    add_section_heading(doc, "3.0 EXISTING CONDITIONS")
    add_mixed_para(doc, [
        ("The inspection covered all accessible external facades and internal areas of the property at ", False),
        (PROJECT["address"], True),
        (". The property was found to be in generally fair condition with some defects present typical "
         "for the age of the structure. Photos and descriptions of condition can be found in the "
         "appendices overleaf.", False),
    ])
    doc.add_paragraph()


def add_photo_section_table(doc, photos: list[dict]):
    """One 2-row table per photo: image row + caption row."""
    content_width = 10204  # 11906 - 2×851 DXA

    for photo in photos:
        photo_path = Path(photo["path"])
        table = _make_borderless_table(doc, 2, 1, [content_width])

        tblPr = table._tbl.find(qn("w:tblPr"))
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tblPr.append(jc)

        img_cell = table.cell(0, 0)
        _set_cell_width(img_cell, content_width)
        img_para = img_cell.paragraphs[0]
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_run = img_para.add_run()
        if photo_path.exists():
            try:
                img_run.add_picture(str(photo_path), width=Inches(4.724))
            except Exception:
                img_para.add_run(f"[Photo: {photo_path.name}]")
        else:
            img_para.add_run(f"[Photo not found: {photo_path.name}]")

        cap_cell = table.cell(1, 0)
        _set_cell_width(cap_cell, content_width)
        _caption_para(cap_cell, photo["number"], photo["description"])

        doc.add_paragraph()


def add_conclusion(doc):
    doc.add_page_break()
    add_section_heading(doc, "4.0 CONCLUSION")
    add_mixed_para(doc, [
        ("The inspection covered all accessible external facades and internal areas of the property at ", False),
        (PROJECT["address"], True),
        (". The property was found to be in generally fair to good condition with some defects present "
         "typical for the age of the structure. No signs of significant structural distress attributable "
         "to the neighbouring development were observed.", False),
    ])
    add_body(doc, (
        f"A total of {PROJECT['total_photos']} photos was taken during our inspection "
        f"({PROJECT['inspection_date']}). A full set of photos can be downloaded via the "
        f"following link: {PROJECT['photo_link']}"
    ))
    add_body(doc, (
        "I am an appropriately qualified person competent in this area. I possess indemnity "
        "insurance to the satisfaction of the client. We trust that this dilapidation report "
        "meets your requirements."
    ))
    doc.add_paragraph()
    doc.add_paragraph()
    _add_signature_table(doc)


def build_building_report(all_photos: list[dict], output_path: str,
                          template_path: str = None, map_paths: dict = None):
    if template_path and Path(template_path).exists():
        doc = Document(template_path)
        body = doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        for child in list(body):
            body.remove(child)
        if sectPr is not None:
            body.append(sectPr)
        print(f"Using template: {Path(template_path).name}")
    else:
        doc = Document()
        _setup_document(doc)

    _add_page_number_footer(doc)

    add_cover_page(doc, cover_photo=map_paths.get("cover") if map_paths else None)
    add_contents(doc)
    add_report_metadata(doc)
    add_preamble(doc)
    add_introduction(doc, map_paths=map_paths)
    add_existing_conditions_intro(doc)
    add_conclusion(doc)

    doc.add_page_break()

    current_appendix = None
    current_facade = None
    section_photos: list[dict] = []

    for photo in all_photos:
        appendix = photo.get("appendix", photo["section"])
        facade = photo.get("facade", "")

        if appendix != current_appendix or facade != current_facade:
            if section_photos:
                add_photo_section_table(doc, section_photos)
                section_photos = []

            if appendix != current_appendix:
                current_appendix = appendix
                current_facade = None
                add_appendix_label(doc, appendix)
                doc.add_paragraph()

            if facade and facade != current_facade:
                current_facade = facade
                add_facade_label(doc, facade)
                doc.add_paragraph()

        section_photos.append(photo)

    if section_photos:
        add_photo_section_table(doc, section_photos)

    doc.save(output_path)
    print(f"\nReport saved to: {output_path}")
